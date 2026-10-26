"""Reusable data cleaning and seven-calendar-day forecasting utilities."""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd

DIMENSIONS = ["State", "District", "Market", "Commodity", "Variety", "Grade"]
DATE = "Arrival_Date"
PRICE = "Modal_Price"
KEYS = DIMENSIONS + [DATE]
ALIASES = {
    "state": "State", "district": "District", "market": "Market",
    "commodity": "Commodity", "variety": "Variety", "grade": "Grade",
    "arrival_date": DATE, "date": DATE, "min_price": "Min_Price",
    "max_price": "Max_Price", "modal_price": PRICE,
    "commodity_code": "Commodity_Code",
}


def parse_dates(values: pd.Series) -> pd.Series:
    """Parse mixed ISO and day-first dates without changing the source values."""
    try:
        return pd.to_datetime(values, format="mixed", dayfirst=True, errors="coerce")
    except (TypeError, ValueError):  # pandas before format="mixed"
        return pd.to_datetime(values, dayfirst=True, errors="coerce")


def read_sources(paths: list[Path]) -> pd.DataFrame:
    """Read CSV/Excel sources; output artifacts should not be added as inputs."""
    frames = []
    for path in sorted({p.resolve() for p in paths}):
        suffix = path.suffix.lower()
        if suffix == ".csv":
            frames.append(pd.read_csv(path, low_memory=False))
        elif suffix == ".xlsx":
            # Stream worksheet rows to keep openpyxl's memory footprint down.
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = workbook.active.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                frames.append(pd.DataFrame.from_records(rows, columns=header))
            finally:
                workbook.close()
        elif suffix == ".xls":
            frames.append(pd.read_excel(path))
        else:
            raise ValueError(f"Unsupported data file: {path.name}")
    if not frames:
        raise FileNotFoundError("No CSV or Excel source files were found.")
    return pd.concat(frames, ignore_index=True, sort=False)


def clean_data(raw: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Normalize, validate, report, and deduplicate records conservatively."""
    df = raw.copy()
    df.columns = [str(c).strip().replace(" ", "_").replace("-", "_").lower() for c in df.columns]
    df = df.rename(columns={k: v for k, v in ALIASES.items() if k in df.columns})
    required = [*DIMENSIONS, DATE, PRICE]
    missing_columns = [c for c in required if c not in df.columns]
    if missing_columns:
        raise ValueError("Missing required columns: " + ", ".join(missing_columns))
    report = {"input_rows": int(len(df)), "missing_columns": missing_columns}
    df[DATE] = parse_dates(df[DATE])
    report["invalid_dates"] = int(df[DATE].isna().sum())
    for col in ["Min_Price", "Max_Price", PRICE]:
        if col in df:
            original_missing = int(df[col].isna().sum())
            df[col] = pd.to_numeric(df[col].astype(str).str.replace(",", "", regex=False), errors="coerce")
            report_name = col.lower()
            report[f"{report_name}_missing_or_invalid"] = int(df[col].isna().sum())
            report[f"{report_name}_zeros"] = int(df[col].eq(0).sum())
            report[f"{report_name}_negative"] = int(df[col].lt(0).sum())
    report["invalid_modal_prices"] = int(df[PRICE].isna().sum())
    report["zero_modal_prices"] = int(df[PRICE].eq(0).sum())
    report["negative_modal_prices"] = int(df[PRICE].lt(0).sum())
    report["missing_modal_prices"] = int(df[PRICE].isna().sum())
    report["missing_categories"] = {col: int(df[col].isna().sum()) for col in DIMENSIONS}
    report["price_distribution"] = {
        key: float(value) for key, value in df[PRICE].quantile([0, .25, .5, .75, 1]).items()
    } if df[PRICE].notna().any() else {}
    for col in DIMENSIONS:
        df[col] = df[col].fillna("Unknown").astype(str).str.strip().replace("", "Unknown")
    unusable = df[DATE].isna() | df[PRICE].isna() | df[PRICE].le(0)
    report["invalid_rows_removed"] = int(unusable.sum())
    df = df.loc[~unusable].copy()
    df[DATE] = df[DATE].dt.normalize()
    report["duplicate_records"] = int(df.duplicated(KEYS, keep="first").sum())
    df = df.drop_duplicates(KEYS, keep="first")
    report["price_relationship_violations"] = 0
    if "Min_Price" in df:
        report["modal_below_min"] = int((df[PRICE] < df["Min_Price"]).sum())
        report["price_relationship_violations"] += report["modal_below_min"]
    if "Max_Price" in df:
        report["modal_above_max"] = int((df[PRICE] > df["Max_Price"]).sum())
        report["price_relationship_violations"] += report["modal_above_max"]
    report.update({
        "retained_rows": int(len(df)),
        "date_min": str(df[DATE].min().date()) if len(df) else None,
        "date_max": str(df[DATE].max().date()) if len(df) else None,
        "commodities": int(df.Commodity.nunique()),
        "markets": int(df.Market.nunique()),
        "date_years": sorted(int(year) for year in df[DATE].dt.year.unique()),
    })
    return df.sort_values(KEYS).reset_index(drop=True), report


def make_features(clean: pd.DataFrame) -> pd.DataFrame:
    """Create exact +7 calendar day targets and strictly prior price features."""
    df = clean.sort_values(KEYS).copy()
    series_keys = DIMENSIONS
    grouped = df.groupby(series_keys, sort=False, dropna=False)[PRICE]
    df["Lag_1"] = grouped.shift(1)
    df["Lag_7"] = grouped.shift(7)
    df["Rolling_Mean_7"] = grouped.transform(lambda s: s.shift(1).rolling(7, min_periods=3).mean())
    df["Rolling_Std_7"] = grouped.transform(lambda s: s.shift(1).rolling(7, min_periods=3).std())
    future = df[series_keys + [DATE, PRICE]].copy()
    future[DATE] = future[DATE] - pd.Timedelta(days=7)
    future = future.rename(columns={PRICE: "Target_Price"})
    df = df.merge(future, on=series_keys + [DATE], how="left", validate="one_to_one")
    df["Target_Date"] = df[DATE] + pd.Timedelta(days=7)
    dt = df[DATE]
    df["Year"], df["Month"], df["Day"], df["Weekday"] = dt.dt.year, dt.dt.month, dt.dt.day, dt.dt.dayofweek
    df["DayOfYear"] = dt.dt.dayofyear
    return df


def chron_split(features: pd.DataFrame, train=0.70, valid=0.15):
    """Split by unique date boundaries, retaining chronological ordering."""
    dates = np.array(sorted(features[DATE].dropna().unique()))
    if len(dates) < 3:
        raise ValueError("At least three distinct observation dates are required.")
    train_end = dates[max(1, int(len(dates) * train)) - 1]
    valid_end = dates[max(2, int(len(dates) * (train + valid))) - 1]
    return (features[DATE] <= train_end,
            (features[DATE] > train_end) & (features[DATE] <= valid_end),
            features[DATE] > valid_end)


def save_clean_report(report: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, indent=2), encoding="utf-8")

